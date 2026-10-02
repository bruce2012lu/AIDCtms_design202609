# -*- coding: utf-8 -*-
"""几何校验：估算文本框溢出与越界。无需 Office。"""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
PT_IN = 72.0


def char_w(ch, size):
    o = ord(ch)
    if o > 0x2E80:                     # CJK 全角
        return size
    if ch in "，。、；：？！（）「」·—…“”":
        return size
    return size * 0.55                 # 拉丁/数字近似


def est_height(text, size, box_w_in, line_spacing=1.15):
    """返回估算所需高度（英寸）"""
    if not text.strip():
        return 0.0
    usable_pt = box_w_in * PT_IN - 6   # 去掉左右内边距
    lines, cur = 1, 0.0
    for ch in text:
        if ch == "\n":
            lines += 1
            cur = 0.0
            continue
        w = char_w(ch, size)
        if cur + w > usable_pt:
            lines += 1
            cur = w
        else:
            cur += w
    return lines * size * 1.32 * line_spacing / PT_IN


def walk(prs):
    problems = []
    sw, sh = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    for idx, slide in enumerate(prs.slides, 1):
        for sp in slide.shapes:
            if not sp.has_text_frame:
                continue
            tf = sp.text_frame
            txt = tf.text
            if not txt.strip():
                continue
            l, t = sp.left / EMU_IN, sp.top / EMU_IN
            w, h = sp.width / EMU_IN, sp.height / EMU_IN
            # 越界
            if l < -0.02 or t < -0.02 or l + w > sw + 0.02 or t + h > sh + 0.02:
                problems.append((idx, "OUT", f"({l:.2f},{t:.2f}) {w:.2f}x{h:.2f}", txt[:34]))
            # 逐段累加
            need = 0.0
            for p in tf.paragraphs:
                ptxt = "".join(r.text for r in p.runs)
                if not ptxt:
                    continue
                size = None
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                        break
                size = size or 18.0
                ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.15
                sa = p.space_after.pt if p.space_after else 0
                need += est_height(ptxt, size, w, ls) + sa / PT_IN
            if need > h + 0.06:
                problems.append((idx, "OVER", f"need {need:.2f} > {h:.2f} (w={w:.2f})", txt[:34]))
    return problems


def main(path):
    prs = Presentation(path)
    probs = walk(prs)
    if not probs:
        print("OK: no overflow / out-of-bounds detected")
        return
    print(f"{len(probs)} issue(s):")
    for idx, kind, detail, txt in probs:
        print(f"  p{idx:02d} {kind:4s} {detail:36s} | {txt}")


if __name__ == "__main__":
    main(sys.argv[1])
