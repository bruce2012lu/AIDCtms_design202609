# -*- coding: utf-8 -*-
"""
v2.2 BP / v3.3 报告 / v3.3 演示版 —— 结构、编码与交叉一致性校验
运行： python verify_v33.py
"""
import os
import re
import sys
from collections import Counter
from html.parser import HTMLParser

HERE = os.path.dirname(os.path.abspath(__file__))
BP_DIR = os.path.dirname(os.path.dirname(HERE))

BP = os.path.join(BP_DIR, "算力冷却系统新业务BP_v2.2_20260906.html")
RP = os.path.join(BP_DIR, "算力冷却商业调研报告_v3.3_20260906.html")
PPTX = os.path.join(BP_DIR, "算力冷却商业调研报告_v3.3_演示版_20260906.pptx")

VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link",
        "meta", "param", "source", "track", "wbr"}

errors = []
warnings = []
oks = []


def err(m):
    errors.append(m)


def warn(m):
    warnings.append(m)


def ok(m):
    oks.append(m)


# ---------------------------------------------------------------- HTML 结构
class Checker(HTMLParser):
    def __init__(self, name):
        super().__init__(convert_charrefs=True)
        self.name = name
        self.stack = []
        self.ids = Counter()
        self.anchors = []
        self.tbl = None          # (ncols_header, rowidx)
        self.in_table = 0
        self.cells = 0
        self.rowspec = []
        self.hdr_cols = None
        self.bad_rows = []
        self.row_line = 0
        self.pending_span = 0

    def handle_starttag(self, tag, attrs):
        d = dict(attrs)
        if "id" in d:
            self.ids[d["id"]] += 1
        if tag == "a" and d.get("href", "").startswith("#"):
            self.anchors.append(d["href"][1:])
        if tag not in VOID:
            self.stack.append((tag, self.getpos()[0]))
        if tag == "table":
            self.in_table += 1
            self.hdr_cols = None
            self.carry = []      # 由 rowspan 带入后续行的占位列数
        if tag == "tr":
            self.cells = sum(n for n, r in getattr(self, "carry", []))
            self.new_spans = []
            self.row_line = self.getpos()[0]
        if tag in ("td", "th"):
            cs = int(d.get("colspan", 1))
            rs = int(d.get("rowspan", 1))
            self.cells += cs
            if rs > 1:
                self.new_spans.append((cs, rs - 1))

    def handle_endtag(self, tag):
        if tag in VOID:
            return
        if not self.stack:
            err(f"[{self.name}] 多余的结束标签 </{tag}>")
            return
        t, ln = self.stack[-1]
        if t != tag:
            err(f"[{self.name}] 标签不匹配：<{t}>（第 {ln} 行）被 </{tag}> 关闭"
                f"（第 {self.getpos()[0]} 行）")
            # 尝试回退
            for i in range(len(self.stack) - 1, -1, -1):
                if self.stack[i][0] == tag:
                    del self.stack[i:]
                    return
            return
        self.stack.pop()
        if tag == "tr" and self.in_table:
            if self.hdr_cols is None:
                self.hdr_cols = self.cells
            elif self.cells != self.hdr_cols:
                self.bad_rows.append((self.row_line, self.cells, self.hdr_cols))
            self.carry = ([(n, r - 1) for n, r in self.carry if r - 1 > 0]
                          + self.new_spans)
            self.new_spans = []
        if tag == "table":
            self.in_table -= 1


def check_html(path):
    name = os.path.basename(path)
    raw = open(path, "rb").read()
    try:
        txt = raw.decode("utf-8")
    except UnicodeDecodeError as e:
        err(f"[{name}] 不是合法 UTF-8：{e}")
        return None
    if raw.startswith(b"\xef\xbb\xbf"):
        warn(f"[{name}] 含 UTF-8 BOM")
    else:
        ok(f"[{name}] 编码 UTF-8 无 BOM")
    if not re.search(r'charset\s*=\s*"?utf-8"?', txt, re.I):
        err(f"[{name}] 缺少 charset=utf-8 声明")
    if "\ufffd" in txt:
        err(f"[{name}] 含替换字符 U+FFFD（乱码）")
    else:
        ok(f"[{name}] 无乱码字符")

    c = Checker(name)
    c.feed(txt)
    if c.stack:
        for t, ln in c.stack:
            err(f"[{name}] 未闭合标签 <{t}>（第 {ln} 行）")
    else:
        ok(f"[{name}] 全部标签闭合")

    dup = [k for k, v in c.ids.items() if v > 1]
    if dup:
        err(f"[{name}] 重复 id：{dup}")
    else:
        ok(f"[{name}] 无重复 id（共 {len(c.ids)} 个）")

    missing = sorted({a for a in c.anchors if a and a not in c.ids and a != "top"})
    # #top 在两份文件中均以 id=\"top\" 存在
    missing = [m for m in missing if m not in c.ids]
    if missing:
        err(f"[{name}] 锚点指向不存在的 id：{missing}")
    else:
        ok(f"[{name}] 全部内部锚点有效（共 {len(set(c.anchors))} 个目标）")

    if c.bad_rows:
        for ln, got, exp in c.bad_rows:
            err(f"[{name}] 表格列数不一致：第 {ln} 行 {got} 列，表头 {exp} 列")
    else:
        ok(f"[{name}] 全部表格列数与表头一致")

    # CSS 类完整性
    defined = set()
    for m in re.finditer(r"<style>(.*?)</style>", txt, re.S):
        for sel in re.finditer(r"\.([A-Za-z][\w-]*)", m.group(1)):
            defined.add(sel.group(1))
    used = set()
    for m in re.finditer(r'class="([^"]+)"', txt):
        used.update(m.group(1).split())
    undef = sorted(used - defined)
    if undef:
        err(f"[{name}] 使用了未定义的 CSS 类：{undef}")
    else:
        ok(f"[{name}] 无未定义 CSS 类（使用 {len(used)} 个）")
    return txt


# ---------------------------------------------------------------- 交叉一致性
# 关键数字：必须在三份材料中一致出现
KEY_HTML = {
    "毛利锚点 10.36%": ["10.36%"],
    "价格上限 684.75 / 715.00 元/kW": ["684.75", "715.00"],
    "同边界成本下界 840 元/kW": ["840"],
    "所需毛利率 27%–37%": ["27%–37%"],
    "国内实证区间 10.36%–23.83%": ["10.36%–23.83%"],
    "敞口倍数 1.76–2.52 / 17.0–24.3": ["1.76–2.52", "17.0–24.3"],
    "责任上限覆盖率 4%–11%": ["4%–11%"],
    "峰值现金 91 万元/MW": ["91 万元/MW"],
    "并行容量 1.21 MW": ["1.21 MW"],
    "年吞吐 1.86 MW": ["1.86 MW"],
    "盈亏平衡 7.80 MW": ["7.80 MW"],
    "达成率 23.8%": ["23.8%"],
    "现金底线 171 万元": ["171"],
    "回款容忍 11.2 周 → 17.1 周": ["11.2 周", "17.1 周"],
    "保护比 49:1": ["49"],
    "受保护支出 245 万元": ["245 万元"],
    "取证成本 <5 万元 / 6–8 周": ["6–8 周"],
    "敞口比 4.71 倍": ["4.71"],
}


def cross_check(bp_txt, rp_txt, ppt_txt):
    both = {"BP v2.2": bp_txt, "报告 v3.3": rp_txt, "PPT v3.3": ppt_txt}
    for label, toks in KEY_HTML.items():
        miss = []
        for docname, doc in both.items():
            if doc is None:
                continue
            for t in toks:
                if t not in doc:
                    miss.append(f"{docname} 缺 {t!r}")
        if miss:
            warn(f"[交叉] {label}：{'；'.join(miss)}")
        else:
            ok(f"[交叉] {label} —— 三份材料一致")


def version_check(bp_txt, rp_txt):
    checks = [
        ("BP v2.2", bp_txt, ["v2.2", "20260906", "2026-09-05"],
         ["v2.1 OEM Operating Model / 20260822</"]),
        ("报告 v3.3", rp_txt, ["v3.3", "20260906"],
         ["v3.2 Evidence-led / 20260822</"]),
    ]
    for name, txt, must, mustnot in checks:
        if txt is None:
            continue
        for m in must:
            if m not in txt:
                err(f"[{name}] 缺少版本标识 {m!r}")
        for m in mustnot:
            if m in txt:
                err(f"[{name}] 仍残留旧版本标识 {m!r}")
        ok(f"[{name}] 版本号与日期检查通过")


# ---------------------------------------------------------------- PPTX
def check_pptx():
    if not os.path.exists(PPTX):
        warn(f"[PPTX] 尚未生成：{PPTX}")
        return None
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(PPTX)
    n = len(prs.slides._sldIdLst)
    ok(f"[PPTX] 可正常打开，共 {n} 页")

    sw, sh = prs.slide_width, prs.slide_height
    texts = []
    imgs = 0
    external = 0
    overflow = []
    bad_font = set()
    for i, sl in enumerate(prs.slides, 1):
        for shp in sl.shapes:
            if shp.shape_type == 13 or shp.__class__.__name__ == "Picture":
                imgs += 1
                try:
                    if shp._element.blip_rId is None:
                        external += 1
                except Exception:
                    pass
            try:
                l, t = shp.left, shp.top
                w, h = shp.width, shp.height
            except Exception:
                continue
            if None in (l, t, w, h):
                continue
            if l < Emu(-9525) or t < Emu(-9525) or l + w > sw + Emu(9525) \
               or t + h > sh + Emu(9525):
                overflow.append((i, shp.shape_type, shp.name,
                                 l / 914400, t / 914400, w / 914400, h / 914400))
            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    for r in p.runs:
                        texts.append(r.text)
                        if r.font.name and r.font.name != "Microsoft YaHei":
                            bad_font.add(r.font.name)
            if shp.has_table:
                for row in shp.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                texts.append(r.text)
                                if r.font.name and r.font.name != "Microsoft YaHei":
                                    bad_font.add(r.font.name)

    body = "\n".join(texts)
    if "\ufffd" in body or "?" * 3 in body:
        err("[PPTX] 疑似中文乱码")
    else:
        ok("[PPTX] 中文无乱码")
    if bad_font:
        warn(f"[PPTX] 出现非 Microsoft YaHei 字体：{sorted(bad_font)}")
    else:
        ok("[PPTX] 全部文本字体为 Microsoft YaHei")
    if overflow:
        for o in overflow:
            err("[PPTX] 第 %d 页形状出框：%s  L%.2f T%.2f W%.2f H%.2f"
                % (o[0], o[2], o[3], o[4], o[5], o[6]))
    else:
        ok("[PPTX] 无形状越出页面边界")
    if external:
        err(f"[PPTX] 存在 {external} 张外链图片")
    else:
        ok(f"[PPTX] {imgs} 张图片全部内嵌")

    import zipfile
    with zipfile.ZipFile(PPTX) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
    ok(f"[PPTX] 打包媒体文件 {len(media)} 个")
    return body


def main():
    print("=" * 78)
    bp = check_html(BP)
    rp = check_html(RP)
    ppt = check_pptx()
    version_check(bp, rp)
    cross_check(bp, rp, ppt)

    print("\n---- OK（%d）----" % len(oks))
    for m in oks:
        print("  OK  ", m)
    if warnings:
        print("\n---- WARN（%d）----" % len(warnings))
        for m in warnings:
            print("  WARN", m)
    if errors:
        print("\n---- ERROR（%d）----" % len(errors))
        for m in errors:
            print("  ERR ", m)
        sys.exit(1)
    print("\n全部检查通过。")


if __name__ == "__main__":
    main()
